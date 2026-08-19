"""
Чтение банковских выписок в формате Word (.docx, а также .doc сохранённый как
OOXML-Word). Внутри — те же выписки, что и в xlsx, просто отрисованные таблицами
Word. Переиспользуем логику reader.py через адаптер GridSheet (DRY).

Каждая таблица документа → GridSheet → find_columns/process_sheet.
Сводные таблицы («Сумма по … счёта»/«Остаток») пропускаются, детальные —
суммируются. Если детальной таблицы дебит/кредит нет (напр. «Справка об
остатках») — возвращаем результат с ошибкой обработки (не исключение).
"""
from app.core.models import AnalysisResult, SheetResult
from app.core.grid import GridSheet
from app.core.reader import find_columns, process_sheet, is_summary_header


def _table_to_rows(table) -> list:
    """Таблица python-docx → 2D-список текстов ячеек."""
    rows = []
    for row in table.rows:
        rows.append([cell.text for cell in row.cells])
    return rows


def parse_docx_file(filepath: str) -> AnalysisResult:
    result = AnalysisResult(filepath=filepath)

    try:
        from docx import Document
        doc = Document(filepath)
    except Exception as e:
        sr = SheetResult(sheet_name="Word")
        sr.error = f"не удалось открыть Word-документ: {e}"
        result.sheets.append(sr)
        return result

    for i, table in enumerate(doc.tables):
        rows2d = _table_to_rows(table)
        if not rows2d:
            continue
        gs = GridSheet(rows2d, title=f"Таблица {i + 1}")

        hrow, _date, dcol, ccol = find_columns(gs)
        if dcol is None:
            continue  # в этой таблице нет столбцов дебит/кредит (реквизиты и т.п.)

        # Сводная таблица итогов («Сумма по … счёта») — не считаем как операции.
        d_text = str(gs.cell(row=hrow, column=dcol).value or "")
        c_text = str(gs.cell(row=hrow, column=ccol).value or "")
        if is_summary_header(d_text, c_text):
            continue

        result.sheets.append(process_sheet(gs))

    if not result.sheets:
        sr = SheetResult(sheet_name="Word")
        sr.error = "таблица с колонками 'дебит'/'кредит' не найдена"
        result.sheets.append(sr)

    return result
